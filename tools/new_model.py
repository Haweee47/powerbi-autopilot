"""Build a semantic model (TMDL) from data you already have: CSV files, an Excel workbook, or a database over ODBC.

Why: the pilots need a model. If your data is already in Power BI, save that report as .pbip and point the agent at it.
If it is not, this writes the model for you, so `new_report.py --model` has something to map onto.

  python tools/new_model.py --name Sales --out out/sales-model --csv data/           # one table per CSV file
  python tools/new_model.py --name Sales --out out/sales-model --excel book.xlsx     # one table per sheet
  python tools/new_model.py --name Sales --out out/sales-model --odbc "Driver={...};..." --table Orders --table Stores
  python tools/new_model.py ... --odbc "<conn>" --query "Orders=SELECT * FROM sales.orders WHERE year >= 2025"
  python tools/new_model.py ... --odbc "<conn>" --table Orders --connector sqlserver --server HOST --database Sales

What it does: reads the column names and a sample of rows, picks a type per column, infers relationships from key columns,
adds a Calendar over the dates it found and one SUM measure per numeric column. Nothing here is a guess you cannot see:
every choice is printed and written as plain TMDL you can edit.

ODBC note: the connection string is written into the model as a parameter, so keep passwords out of it (use a DSN or
integrated security). Credentials are never written by this script. The first refresh in Desktop asks how to sign in.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import re
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
M_TYPE = {"string": "type text", "int64": "Int64.Type", "double": "type number", "dateTime": "type datetime", "boolean": "type logical"}
KEY_HINT = re.compile(r"(^|[ _])(id|key|code|no|번호|코드)$", re.I)
DATE_FORMATS = ("%Y-%m-%d", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%m/%d/%Y", "%d.%m.%Y")
# Native connectors this script can write. Columns still come from an ODBC profile, so the shape is never guessed.
CONNECTORS = {
    "odbc": None,
    "sqlserver": 'Sql.Database(Server, Database){{[Schema="{schema}", Item="{item}"]}}[Data]',
    "mysql": 'MySQL.Database(Server, Database){{[Schema="{schema}", Item="{item}"]}}[Data]',
    "postgres": 'PostgreSQL.Database(Server, Database){{[Schema="{schema}", Item="{item}"]}}[Data]',
    "redshift": 'AmazonRedshift.Database(Server, Database){{[Schema="{schema}", Item="{item}"]}}[Data]',
    "oracle": 'Oracle.Database(Server){{[Schema="{schema}", Item="{item}"]}}[Data]',
}


def q(name: str) -> str:
    return f"'{name}'" if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", name) else name


# ---- reading the source -------------------------------------------------------------------------------------------

def guess_type(values: list[str]) -> str:
    """The narrowest type every non-empty value fits. Strings that would lose a leading zero stay text."""
    vals = [v for v in values if v not in ("", None)]
    if not vals:
        return "string"
    if all(str(v).strip().lower() in ("true", "false") for v in vals):
        return "boolean"
    if all(re.fullmatch(r"-?\d{1,18}", str(v).strip()) for v in vals):
        return "string" if any(re.fullmatch(r"0\d+", str(v).strip()) for v in vals) else "int64"
    if all(re.fullmatch(r"-?\d*\.?\d+([eE][-+]?\d+)?", str(v).strip()) for v in vals):
        return "double"
    for fmt in DATE_FORMATS:
        try:
            for v in vals:
                dt.datetime.strptime(str(v).strip(), fmt)
            return "dateTime"
        except ValueError:
            continue
    return "string"


def from_csv(path: Path, sample: int) -> tuple[list[str], list[list], dict]:
    with open(path, encoding="utf-8-sig", newline="") as f:
        head = f.read(4096)
        f.seek(0)
        delim = csv.Sniffer().sniff(head, delimiters=",;\t|").delimiter if head.strip() else ","
        r = csv.reader(f, delimiter=delim)
        header = next(r, [])
        rows = [row for _, row in zip(range(sample), r)]
    return header, rows, {"delimiter": delim}


def from_excel(path: Path, sheet: str, sample: int) -> tuple[list[str], list[list]]:
    from openpyxl import load_workbook  # only needed for --excel
    ws = load_workbook(path, read_only=True, data_only=True)[sheet]
    it = ws.iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else "" for c in next(it, [])]
    rows = [list(row) for _, row in zip(range(sample), it)]
    return header, rows


def excel_sheets(path: Path) -> list[str]:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    return [ws.title for ws in wb.worksheets if ws.max_row and ws.max_row > 1]


PS_ODBC = """$ErrorActionPreference = 'Stop'
$c = New-Object System.Data.Odbc.OdbcConnection($env:AP_CONN); $c.Open()
$cmd = $c.CreateCommand(); $cmd.CommandText = $env:AP_SQL
$r = $cmd.ExecuteReader(); $cols = @(); $types = @()
for ($i = 0; $i -lt $r.FieldCount; $i++) { $cols += $r.GetName($i); $types += $r.GetFieldType($i).Name }
$rows = @(); $n = 0
while ($r.Read() -and $n -lt [int]$env:AP_N) { $v = @(); for ($i = 0; $i -lt $r.FieldCount; $i++) { $v += [string]$r.GetValue($i) }; $rows += ,$v; $n++ }
$c.Close(); @{ columns = $cols; types = $types; rows = $rows } | ConvertTo-Json -Depth 4 -Compress
"""


# What the driver reports a column as, rather than what its text happens to look like
NET_TYPE = {"DateTime": "dateTime", "DateTimeOffset": "dateTime", "Boolean": "boolean", "Byte": "int64", "Int16": "int64",
            "Int32": "int64", "Int64": "int64", "Single": "double", "Double": "double", "Decimal": "double"}


def from_odbc(conn: str, sql: str, sample: int) -> tuple[list[str], list[str], list[list]]:
    """Read column names, the driver's own types and a sample, using .NET's ODBC client so nothing extra has to be installed."""
    p = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", PS_ODBC],
                       capture_output=True, text=True, env={**__import__("os").environ,
                                                            "AP_CONN": conn, "AP_SQL": sql, "AP_N": str(sample)})
    if p.returncode or not p.stdout.strip():
        sys.exit(f"ODBC read failed for: {sql}\n{p.stderr.strip()[:1200]}")
    data = json.loads(p.stdout)
    rows = data.get("rows") or []
    return data["columns"], data.get("types") or [], [r if isinstance(r, list) else [r] for r in rows]


# ---- shaping the model --------------------------------------------------------------------------------------------

def columns_of(header: list[str], rows: list[list], declared: list[str] | None = None) -> list[tuple[str, str]]:
    """`declared` is the source's own type per column (ODBC). Where the source states a type, it wins over sniffing."""
    out = []
    for i, name in enumerate(header):
        name = (str(name) if name is not None else "").strip() or f"Column {i + 1}"
        values = [row[i] if i < len(row) else None for row in rows]
        if declared and i < len(declared) and declared[i] in NET_TYPE:
            out.append((name, NET_TYPE[declared[i]]))
            continue
        if values and all(isinstance(v, (dt.datetime, dt.date)) for v in values if v is not None):
            out.append((name, "dateTime"))
        elif values and all(isinstance(v, bool) for v in values if v is not None):
            out.append((name, "boolean"))
        elif values and all(isinstance(v, int) and not isinstance(v, bool) for v in values if v is not None):
            out.append((name, "int64"))
        elif values and all(isinstance(v, float) for v in values if v is not None):
            out.append((name, "double"))
        else:
            out.append((name, guess_type([v for v in values])))
    return out


def relationships(tables: dict) -> list[tuple[str, str, str]]:
    """A column that repeats in another table and is unique there is a key: many-to-one, the unique side is the lookup."""
    rels, seen = [], set()
    for fact, f in tables.items():
        for col, _ in f["columns"]:
            for dim, d in tables.items():
                if dim == fact or col not in [c for c, _ in d["columns"]]:
                    continue
                i = [c for c, _ in d["columns"]].index(col)
                vals = [str(r[i]) for r in d["rows"] if i < len(r) and r[i] not in (None, "")]
                unique = vals and len(set(vals)) == len(vals)
                if not unique or not (KEY_HINT.search(col) or len(d["rows"]) < len(f["rows"])):
                    continue
                name = re.sub(r"\W", "", f"{fact}_{dim}_{col}")
                if (fact, col) in seen:
                    continue
                seen.add((fact, col))
                rels.append((name, f"{q(fact)}.{q(col)}", f"{q(dim)}.{q(col)}"))
    return rels


def measures(tables: dict, rels: list) -> list[tuple[str, str, str]]:
    """One SUM per numeric column that is not a key, plus a row count for the biggest table. A starting point to edit."""
    keys = {a.split(".", 1)[1].strip("'") for _, a, _ in rels} | {b.split(".", 1)[1].strip("'") for _, _, b in rels}
    out = []
    for t, info in tables.items():
        for col, dtype in info["columns"]:
            if dtype in ("int64", "double") and col not in keys and not KEY_HINT.search(col):
                out.append((f"Total {col}", f"SUM ( {q(t)}[{col}] )", "#,0.##"))
    if tables:
        big = max(tables, key=lambda t: len(tables[t]["rows"]))
        out.insert(0, (f"{big} Rows", f"COUNTROWS ( {q(big)} )", "#,0"))
    return out


# ---- writing TMDL -------------------------------------------------------------------------------------------------

def table_tmdl(name: str, cols: list[tuple[str, str]], source: str) -> str:
    lines = [f"table {q(name)}\n"]
    for col, dtype in cols:
        lines += [f"\tcolumn {q(col)}", f"\t\tdataType: {dtype}",
                  f"\t\tsummarizeBy: {'sum' if dtype in ('int64', 'double') else 'none'}"]
        if dtype == "dateTime":
            lines.append("\t\tformatString: yyyy-mm-dd")
        lines += [f"\t\tsourceColumn: {col}\n"]
    lines.append(f"\tpartition {q(name)} = m\n\t\tmode: import\n\t\tsource =\n{source}\n")
    return "\n".join(lines)


def m_source(kind: str, name: str, cols: list[tuple[str, str]], opt: dict) -> str:
    """The M for one table. The types come from the profile, so a refresh cannot silently change them."""
    types = ", ".join(f'{{"{c}", {M_TYPE[t]}}}' for c, t in cols)
    i = " " * 16
    if kind == "csv":
        read = (f'Csv.Document(File.Contents(SourcePath & "\\{opt["file"]}"), '
                f'[Delimiter = "{opt["delimiter"]}", Columns = {len(cols)}, Encoding = 65001, QuoteStyle = QuoteStyle.Csv])')
        steps = [f"Source = {read},", "Promoted = Table.PromoteHeaders(Source, [PromoteAllScalars = true]),"]
        last = "Promoted"
    elif kind == "excel":
        steps = ["Book = Excel.Workbook(File.Contents(SourcePath), null, true),",
                 f'Source = Book{{[Item = "{opt["sheet"]}", Kind = "Sheet"]}}[Data],',
                 "Promoted = Table.PromoteHeaders(Source, [PromoteAllScalars = true]),"]
        last = "Promoted"
    elif opt.get("connector", "odbc") == "odbc":
        steps = [f'Source = Odbc.Query(ConnectionString, "{opt["sql"]}"),']
        last = "Source"
    else:
        shape = CONNECTORS[opt["connector"]].format(schema=opt.get("schema", "dbo"), item=opt["item"])
        steps = [f"Source = {shape},"]
        last = "Source"
    body = "\n".join(f"{i}    {s}" for s in steps)
    return (f"{i}let\n{body}\n{i}    Typed = Table.TransformColumnTypes({last}, {{{types}}}, \"en-US\")\n"
            f"{i}in\n{i}    Typed")


CALENDAR = ('ADDCOLUMNS ( CALENDARAUTO (), "Year", YEAR ( [Date] ), "Month No", MONTH ( [Date] ), '
            '"Month", FORMAT ( [Date], "MMM", "en-US" ), "Quarter", "Q" & QUARTER ( [Date] ), '
            '"Week Start", [Date] - WEEKDAY ( [Date], 3 ) )')
CALENDAR_COLS = [("Date", ["isUnique", "formatString: yyyy-mm-dd", "summarizeBy: none"]), ("Year", ["formatString: 0", "summarizeBy: none"]),
                 ("Month No", ["isHidden", "summarizeBy: none"]), ("Month", ["summarizeBy: none", "sortByColumn: 'Month No'"]),
                 ("Quarter", ["summarizeBy: none"]), ("Week Start", ["formatString: mm-dd", "summarizeBy: none"])]


def calendar_tmdl() -> str:
    lines = ["table Calendar\n\tdataCategory: Time\n"]
    for col, extra in CALENDAR_COLS:
        lines += [f"\tcolumn {q(col)}"] + [f"\t\t{e}" for e in extra] + ["\t\tisNameInferred", f"\t\tsourceColumn: [{col}]\n"]
    lines.append(f"\tpartition Calendar = calculated\n\t\tmode: import\n\t\tsource = {CALENDAR}\n")
    return "\n".join(lines)


def metrics_tmdl(ms: list[tuple[str, str, str]]) -> str:
    lines = ["/// Measure-only table: every metric lives here, the way the pilots expect.", "table Metrics\n"]
    for name, dax, fmt in ms:
        lines += [f"\tmeasure {q(name)} = {dax}", f"\t\tformatString: {fmt}", ""]
    lines += ["\tcolumn _\n\t\tisHidden\n\t\tisNameInferred\n\t\tsourceColumn: [_]\n",
              "\tpartition Metrics = calculated\n\t\tmode: import\n\t\tsource = ROW ( \"_\", BLANK () )\n"]
    return "\n".join(lines)


def write_model(out: Path, name: str, tables: dict, rels: list, ms: list, params: list[tuple[str, str]], calendar: bool) -> None:
    (out / "tables").mkdir(parents=True, exist_ok=True)
    for f in (out / "tables").glob("*.tmdl"):
        f.unlink()
    names = list(tables) + (["Calendar"] if calendar else []) + (["Metrics"] if ms else [])
    (out / "database.tmdl").write_text("database\n\tcompatibilityLevel: 1601\n\tcompatibilityMode: powerBI\n", encoding="utf-8")
    (out / "model.tmdl").write_text(
        f"/// {name}: built by tools/new_model.py from the source below. Edit freely - this is a starting point.\n"
        "model Model\n\tculture: en-US\n\tdefaultPowerBIDataSourceVersion: powerBI_V3\n\tdiscourageImplicitMeasures\n"
        "\tsourceQueryCulture: en-US\n\tdataAccessOptions\n\t\tlegacyRedirects\n\t\treturnErrorValuesAsNull\n\n"
        "annotation __PBI_TimeIntelligenceEnabled = 0\n\n" + "".join(f"ref table {q(t)}\n" for t in names), encoding="utf-8")
    if params:
        (out / "expressions.tmdl").write_text("".join(
            f"/// Change this one value to move the model to another machine or environment.\n"
            f'expression {k} = "{v}" meta [IsParameterQuery=true, Type="Text", IsParameterQueryRequired=true]\n\n'
            for k, v in params), encoding="utf-8")
    (out / "relationships.tmdl").write_text("".join(
        f"relationship {n}\n\tfromColumn: {a}\n\ttoColumn: {b}\n\n" for n, a, b in rels), encoding="utf-8")
    for t, info in tables.items():
        (out / "tables" / f"{t}.tmdl").write_text(table_tmdl(t, info["columns"], info["source"]), encoding="utf-8")
    if calendar:
        (out / "tables" / "Calendar.tmdl").write_text(calendar_tmdl(), encoding="utf-8")
    if ms:
        (out / "tables" / "Metrics.tmdl").write_text(metrics_tmdl(ms), encoding="utf-8")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--name", required=True)
    ap.add_argument("--out", required=True, help="where the .SemanticModel folder content goes")
    ap.add_argument("--csv", nargs="+", help="CSV files, or a folder of them")
    ap.add_argument("--excel", help="workbook: one table per sheet")
    ap.add_argument("--sheet", action="append", help="only these sheets (repeatable)")
    ap.add_argument("--odbc", help="ODBC connection string, without a password")
    ap.add_argument("--table", action="append", default=[], help="table to read over ODBC (repeatable), e.g. dbo.Orders")
    ap.add_argument("--query", action="append", default=[], help='Name=SELECT ... (repeatable)')
    ap.add_argument("--connector", choices=sorted(CONNECTORS), default="odbc",
                    help="M connector to write. Columns are always profiled over ODBC; the native ones need --server/--database")
    ap.add_argument("--server"), ap.add_argument("--database")
    ap.add_argument("--sample", type=int, default=2000, help="rows read to pick column types (default 2000)")
    ap.add_argument("--no-calendar", action="store_true"), ap.add_argument("--no-measures", action="store_true")
    a = ap.parse_args()
    if not (a.csv or a.excel or a.odbc):
        ap.error("give one of --csv, --excel or --odbc")
    if a.connector != "odbc" and not (a.server and a.database or a.connector == "oracle" and a.server):
        ap.error(f"--connector {a.connector} needs --server and --database")

    out, tables, params = Path(a.out), {}, []
    if a.csv:
        paths = [p for arg in a.csv for p in (sorted(Path(arg).glob("*.csv")) if Path(arg).is_dir() else [Path(arg)])]
        if not paths:
            sys.exit("no CSV files found")
        params.append(("SourcePath", str(paths[0].resolve().parent)))
        for p in paths:
            header, rows, opt = from_csv(p, a.sample)
            cols = columns_of(header, rows)
            tables[p.stem] = {"columns": cols, "rows": rows,
                              "source": m_source("csv", p.stem, cols, {"file": p.name, **opt})}
    elif a.excel:
        book = Path(a.excel).resolve()
        params.append(("SourcePath", str(book)))
        for sheet in (a.sheet or excel_sheets(book)):
            header, rows = from_excel(book, sheet, a.sample)
            cols = columns_of(header, rows)
            tables[sheet] = {"columns": cols, "rows": rows, "source": m_source("excel", sheet, cols, {"sheet": sheet})}
    else:
        params.append(("ConnectionString", a.odbc))
        if a.connector != "odbc":
            params += [("Server", a.server or ""), ("Database", a.database or "")]
        jobs = [(t.split(".")[-1], f"SELECT * FROM {t}", t) for t in a.table]
        jobs += [(qq.split("=", 1)[0], qq.split("=", 1)[1], None) for qq in a.query if "=" in qq]
        if not jobs:
            sys.exit("give --table or --query for the ODBC source")
        for name, sql, full in jobs:
            header, declared, rows = from_odbc(a.odbc, sql, a.sample)
            cols = columns_of(header, rows, declared)
            schema, item = (full.split(".", 1) if full and "." in full else ("dbo", full or name))
            tables[name] = {"columns": cols, "rows": rows,
                            "source": m_source("db", name, cols, {"sql": sql.replace('"', '""'), "connector": a.connector,
                                                                  "schema": schema, "item": item})}

    rels = relationships(tables)
    ms = [] if a.no_measures else measures(tables, rels)
    has_date = any(t == "dateTime" for info in tables.values() for _, t in info["columns"])
    calendar = has_date and not a.no_calendar
    if calendar:  # one date column per table joins the calendar, so every table can be filtered by period
        for t, info in tables.items():
            first = next((c for c, dtp in info["columns"] if dtp == "dateTime"), None)
            if first:
                rels.append((re.sub(r"\W", "", f"{t}_Calendar_{first}"), f"{q(t)}.{q(first)}", "Calendar.Date"))
    write_model(out, a.name, tables, rels, ms, params, calendar)
    print(f"{out}: {len(tables) + calendar + bool(ms)} tables, {sum(len(i['columns']) for i in tables.values())} columns, "
          f"{len(rels)} relationships, {len(ms)} measures")
    for t, info in tables.items():
        print(f"  {t}: " + ", ".join(f"{c} ({d})" for c, d in info["columns"][:6]) + (" ..." if len(info["columns"]) > 6 else ""))
    for n, x, y in rels:
        print(f"  relationship {x} -> {y}")
    print("Next: python tools/new_report.py --purpose <dashboard|table|matrix|deepdive|fulfillment> "
          f"--model {out} --name {a.name}")


if __name__ == "__main__":
    main()
