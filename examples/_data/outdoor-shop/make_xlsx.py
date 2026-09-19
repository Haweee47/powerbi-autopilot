"""Put the outdoor-shop CSVs into one Excel workbook, a sheet per table.

Why: many people keep their data in a workbook, not in a folder of CSVs. This makes that case reproducible
for `tools/new_model.py --excel` without committing a binary file.

Usage: python examples/_data/outdoor-shop/make_xlsx.py [outdoor-shop.xlsx]
"""
import csv
import datetime as dt
import sys
from pathlib import Path

from openpyxl import Workbook

HERE = Path(__file__).resolve().parent


def value(text: str):
    """Write dates and numbers as dates and numbers, so the workbook carries types the way a real one does."""
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return dt.datetime.strptime(text, fmt)
        except ValueError:
            pass
    try:
        return int(text) if text.lstrip("-").isdigit() else float(text)
    except ValueError:
        return text


def main() -> None:
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "outdoor-shop.xlsx"
    wb = Workbook()
    wb.remove(wb.active)
    for path in sorted(HERE.glob("*.csv")):
        ws = wb.create_sheet(path.stem)
        with open(path, encoding="utf-8-sig", newline="") as f:
            for i, row in enumerate(csv.reader(f)):
                ws.append(row if i == 0 else [value(v) for v in row])
    wb.save(out)
    print(f"{out}: {', '.join(wb.sheetnames)}")


if __name__ == "__main__":
    main()
